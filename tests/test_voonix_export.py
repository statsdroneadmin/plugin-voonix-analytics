"""
tests/test_voonix_export.py — offline tests for src/voonix_analytics_export.py.

Every writer is round-tripped: the file is written, then read back with an
independent reader (zipfile/csv/json/sqlite3, plus openpyxl when installed).

Run:  python3 tests/test_voonix_export.py
"""

from __future__ import annotations

import csv
import io
import json
import os
import sqlite3
import sys
import tempfile
import traceback
import zipfile
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import voonix_analytics_export as ex  # noqa: E402
import voonix_analytics_mappers as mp  # noqa: E402

INFO = {"exported_at": "2026-09-15T12:00:00+00:00", "version": "0.2.0", "format": "test",
        "start": None, "end": None, "include_raw": True}
TRICKY = 'Casino "Hat", no=\'x\'\nline2 <b>&</b>'


def sample_tables() -> list[ex.ExportTable]:
    return [
        ex.ExportTable(
            name="voonix_sites", label="Sites",
            columns=[("id", "bigint"), ("name", "text"), ("raw", "jsonb")], pk=["id"],
            rows=iter([(1, TRICKY, {"a": 1, "b": "ø"}), (2, None, None)]),
        ),
        ex.ExportTable(
            name="voonix_advertiser_earnings", label="Advertiser earnings",
            columns=[("period_month", "date"), ("row_no", "integer"), ("date", "date"),
                     ("rev_income", "numeric"), ("paused", "boolean"), ("synced_at", "timestamptz")],
            pk=["period_month", "row_no"],
            rows=iter([
                (date(2026, 9, 1), 0, date(2026, 9, 2), Decimal("3071.0010"), True,
                 datetime(2026, 9, 15, 12, 0, tzinfo=timezone.utc)),
                (date(2026, 9, 1), 1, None, Decimal("-5"), False, None),
            ]),
        ),
    ]


def tmp(suffix: str) -> str:
    fd, path = tempfile.mkstemp(suffix=suffix)
    os.close(fd)
    return path


# ── Schema ───────────────────────────────────────────────────────────────────


def test_schema_matches_mappers_and_specs():
    schema = ex.load_schema()
    for table, cols in mp.TABLE_COLUMNS.items():
        names = [c for c, _ in schema[table].columns]
        assert not set(cols) - set(names), f"{table}: {set(cols) - set(names)} missing from migration parse"
    assert schema["voonix_advertiser_earnings"].pk == ["period_month", "row_no"]
    assert schema["voonix_advertisers"].pk == ["id"]
    assert schema["voonix_sync_state"].pk == ["key"]
    assert ("commission", "numeric") in schema["voonix_advertiser_earnings"].columns
    assert ("key", "text") in schema["voonix_campaigns"].columns
    for spec in ex.EXPORT_TABLES:
        cols = dict(schema[spec["table"]].columns)
        assert all(rc in cols for rc in spec["raw_cols"]), spec["table"]
        if spec.get("date_col"):
            assert spec["date_col"] in cols, spec["table"]
    assert {s["table"] for s in ex.EXPORT_TABLES} == set(schema), "every migrated table must be exportable"


def test_select_tables_and_query():
    assert [t["table"] for t in ex.select_tables("")] == [t["table"] for t in ex.EXPORT_TABLES if t["default"]]
    assert [t["table"] for t in ex.select_tables("voonix_sites,voonix_payouts")] == ["voonix_payouts", "voonix_sites"]
    try:
        ex.select_tables("voonix_sites; DROP TABLE x")
    except ValueError:
        pass
    else:
        raise AssertionError("unknown table must raise")

    schema = ex.load_schema()
    spec = ex.select_tables("voonix_advertiser_earnings")[0]
    q = ex.table_query(spec, schema[spec["table"]], date(2026, 1, 15), date(2026, 3, 31), include_raw=False)
    assert "raw" not in [c for c, _ in q["columns"]] and '"raw"' not in q["select_sql"]
    assert q["params"] == [date(2026, 1, 15), date(2026, 3, 31)]
    assert q["select_sql"].endswith('ORDER BY "period_month", "row_no"')

    payouts = ex.select_tables("voonix_payouts")[0]
    q = ex.table_query(payouts, schema["voonix_payouts"], date(2026, 1, 15), None, include_raw=True)
    assert q["params"] == [date(2026, 1, 1)] and "raw" in [c for c, _ in q["columns"]]

    log = ex.select_tables("voonix_write_log")[0]
    q = ex.table_query(log, schema["voonix_write_log"], date(2026, 1, 1), date(2026, 1, 31), include_raw=False)
    assert '"created_at" < (%s::date + 1)' in q["select_sql"]
    assert "request" not in [c for c, _ in q["columns"]]

    advertisers = ex.select_tables("voonix_advertisers")[0]
    q = ex.table_query(advertisers, schema["voonix_advertisers"], date(2026, 1, 1), None, include_raw=False)
    assert q["params"] == [] and "WHERE" not in q["select_sql"], "account tables ignore the date range"


# ── Literals ─────────────────────────────────────────────────────────────────


def test_sql_literal():
    assert ex.sql_literal(None) == "NULL"
    assert ex.sql_literal(True) == "TRUE"
    assert ex.sql_literal(Decimal("3071.00100")) == "3071.001"
    assert ex.sql_literal(Decimal("-5")) == "-5"
    assert ex.sql_literal("it's") == "'it''s'"
    assert ex.sql_literal("a\\b") == "'a\\b'"
    assert ex.sql_literal("nul\x00byte") == "'nulbyte'"
    assert ex.sql_literal({"k": "v'"}) == "'{\"k\": \"v''\"}'"
    assert ex.sql_literal(date(2026, 9, 1)) == "'2026-09-01'"


def test_csv_safe():
    assert ex.csv_safe("=HYPERLINK(1)") == "'=HYPERLINK(1)"
    assert ex.csv_safe("-5") == "'-5"
    assert ex.csv_safe(Decimal("-5")) == "-5", "numbers are never escaped"
    assert ex.csv_safe("Betway") == "Betway"


# ── Writers ──────────────────────────────────────────────────────────────────


def test_csv_zip_roundtrip():
    path = tmp(".zip")
    counts = ex.write_csv_zip(path, sample_tables(), INFO)
    assert counts == {"voonix_sites": 2, "voonix_advertiser_earnings": 2}
    with zipfile.ZipFile(path) as zf:
        assert set(zf.namelist()) == {"sites.csv", "advertiser_earnings.csv", "README.txt"}
        rows = list(csv.reader(io.StringIO(zf.read("sites.csv").decode("utf-8-sig"))))
        assert rows[0] == ["id", "name", "raw"] and rows[1][1] == TRICKY
        assert json.loads(rows[1][2]) == {"a": 1, "b": "ø"} and rows[2] == ["2", "", ""]
        earn = list(csv.reader(io.StringIO(zf.read("advertiser_earnings.csv").decode("utf-8-sig"))))
        assert earn[1] == ["2026-09-01", "0", "2026-09-02", "3071.001", "true", "2026-09-15T12:00:00+00:00"]
        assert "sites: 2" in zf.read("README.txt").decode()
    os.unlink(path)


def test_jsonl_zip_roundtrip():
    path = tmp(".zip")
    ex.write_jsonl_zip(path, sample_tables(), INFO)
    with zipfile.ZipFile(path) as zf:
        lines = zf.read("advertiser_earnings.jsonl").decode().splitlines()
        first = json.loads(lines[0])
        assert first["rev_income"] == 3071.001 and first["paused"] is True and first["date"] == "2026-09-02"
        assert json.loads(lines[1])["rev_income"] == -5
        site = json.loads(zf.read("sites.jsonl").decode().splitlines()[0])
        assert site["name"] == TRICKY and site["raw"] == {"a": 1, "b": "ø"}
    os.unlink(path)


def test_sqlite_roundtrip():
    path = tmp(".sqlite")
    ex.write_sqlite(path, sample_tables(), INFO)
    con = sqlite3.connect(path)
    assert con.execute('SELECT name, raw FROM "voonix_sites" WHERE id = 1').fetchone() == (TRICKY, '{"a": 1, "b": "ø"}')
    assert con.execute('SELECT rev_income, paused FROM "voonix_advertiser_earnings" ORDER BY row_no').fetchall() == [
        (3071.001, 1), (-5, 0)]
    counts = json.loads(con.execute("SELECT value FROM _export_info WHERE key = 'row_counts'").fetchone()[0])
    assert counts["voonix_sites"] == 2
    try:
        con.execute('INSERT INTO "voonix_sites" (id) VALUES (1)')
    except sqlite3.IntegrityError:
        pass
    else:
        raise AssertionError("primary key must be carried into SQLite")
    con.close()
    os.unlink(path)


def test_sql_dump_executes():
    path = tmp(".zip")
    ex.write_sql_zip(path, sample_tables(), INFO, batch=1)
    with zipfile.ZipFile(path) as zf:
        dump = zf.read("voonix-export.sql").decode()
    assert dump.count("INSERT INTO") == 4 and "ON CONFLICT DO NOTHING" in dump and dump.rstrip().endswith("COMMIT;")
    # SQLite understands this INSERT/CREATE dialect apart from the SET lines;
    # the real PostgreSQL load is checked separately against a throwaway cluster.
    con = sqlite3.connect(":memory:")
    con.executescript("\n".join(line for line in dump.splitlines() if not line.startswith("SET ")))
    assert con.execute('SELECT name FROM "voonix_sites" WHERE id = 1').fetchone()[0] == TRICKY
    assert con.execute('SELECT COUNT(*) FROM "voonix_advertiser_earnings"').fetchone()[0] == 2
    os.unlink(path)


def _sheet_rows(zf: zipfile.ZipFile, n: int) -> list[list[str]]:
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    root = ET.fromstring(zf.read(f"xl/worksheets/sheet{n}.xml"))
    out = []
    for row in root.findall(".//m:row", ns):
        cells = []
        for c in row.findall("m:c", ns):
            t = c.find("m:is/m:t", ns)
            v = c.find("m:v", ns)
            cells.append(t.text if t is not None else (v.text if v is not None else ""))
        out.append(cells)
    return out


def test_xlsx_structure_and_values():
    path = tmp(".xlsx")
    ex.write_xlsx(path, sample_tables(), INFO)
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        assert {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml", "xl/styles.xml",
                "xl/_rels/workbook.xml.rels", "xl/worksheets/sheet3.xml"} <= names
        for name in names:
            if name.endswith(".xml") or name.endswith(".rels"):
                ET.fromstring(zf.read(name))  # well-formed
        sites = _sheet_rows(zf, 1)
        assert sites[0] == ["id", "name", "raw"] and sites[1][1] == TRICKY
        assert b'sheet name="Sites"' in zf.read("xl/workbook.xml")
    try:
        import openpyxl  # dev-only verification with an independent reader
    except ImportError:
        print("     (openpyxl not installed — skipped independent Excel read)")
    else:
        wb = openpyxl.load_workbook(path, read_only=True)
        assert wb.sheetnames == ["Sites", "Advertiser earnings", "Export info"]
        rows = list(wb["Advertiser earnings"].iter_rows(values_only=True))
        assert rows[1][3] == 3071.001 and rows[1][4] is True and rows[2][3] == -5
        assert list(wb["Sites"].iter_rows(values_only=True))[1][1] == TRICKY
        wb.close()
    os.unlink(path)


def test_xlsx_splits_sheets_at_row_limit():
    original = ex.XLSX_MAX_ROWS
    ex.XLSX_MAX_ROWS = 3  # header + 2 data rows per sheet
    try:
        path = tmp(".xlsx")
        table = ex.ExportTable("voonix_sites", "Sites", [("id", "bigint")], ["id"], iter([(i,) for i in range(5)]))
        counts = ex.write_xlsx(path, [table], info=None)
        assert counts == {"voonix_sites": 5}
        with zipfile.ZipFile(path) as zf:
            wb = zf.read("xl/workbook.xml").decode()
            assert 'name="Sites"' in wb and 'name="Sites 2"' in wb and 'name="Sites 3"' in wb
            assert [r[0] for r in _sheet_rows(zf, 3)] == ["id", "4"]
        os.unlink(path)
    finally:
        ex.XLSX_MAX_ROWS = original


def test_report_files():
    headers = ["Advertiser", "Commission"]
    rows = [("=cmd()", Decimal("10.50")), ("Betway", Decimal("-3"))]
    path = tmp(".csv")
    assert ex.write_report_file("csv", path, "Advertiser earnings", headers, iter(rows)) == 2
    data = list(csv.reader(open(path, encoding="utf-8-sig")))
    assert data == [headers, ["'=cmd()", "10.5"], ["Betway", "-3"]]
    ex.write_report_file("json", path, "Advertiser earnings", headers, iter(rows))
    assert json.load(open(path)) == [{"Advertiser": "=cmd()", "Commission": 10.5}, {"Advertiser": "Betway", "Commission": -3}]
    ex.write_report_file("json", path, "x", headers, iter([]))
    assert json.load(open(path)) == []
    xlsx = tmp(".xlsx")
    assert ex.write_report_file("xlsx", xlsx, "Advertiser earnings", headers, iter(rows)) == 2
    with zipfile.ZipFile(xlsx) as zf:
        assert _sheet_rows(zf, 1)[1] == ["=cmd()", "10.5"], "xlsx text cells are never formulas"
    os.unlink(path)
    os.unlink(xlsx)


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    failed = 0
    for name, fn in tests:
        try:
            fn()
            print(f"ok   {name}")
        except Exception:
            failed += 1
            print(f"FAIL {name}")
            traceback.print_exc()
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    sys.exit(1 if failed else 0)
